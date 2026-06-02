#!/usr/bin/env python3
"""Fill the Paper8 MSE workbook with V5/Scheme6 averaged results.

The 5.26 workbook is the template. This script leaves it unchanged and writes a
new workbook with V5 columns, V5 comparison sheets, and an all-variant Scheme6
sheet sourced from ``scheme6_collected_results_averaged.csv``.
"""

from __future__ import annotations

import argparse
import csv
from copy import copy
from pathlib import Path

try:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ModuleNotFoundError as exc:
    raise SystemExit(
        "fill_scheme6_excel.py requires openpyxl. "
        "Install project dependencies from Spectral-master/requirements.txt."
    ) from exc


SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = SCRIPT_DIR.parents[1]

DEFAULT_TEMPLATE = "result/5.26_paper8_experiment/final_tables/paper8_mse_improvement_analysis.xlsx"
DEFAULT_SCHEME6_AVERAGED = "Spectral-master/results/summary/scheme6_collected_results_averaged.csv"
DEFAULT_OUTPUT = "Spectral-master/results/summary/paper8_mse_improvement_analysis_scheme6.xlsx"


def project_path(value: str | Path) -> Path:
    path = Path(value)
    if path.is_absolute():
        return path
    return PROJECT_ROOT / path


def normalize_dataset(value) -> str:
    text = str(value).strip()
    key = text.lower().replace(" ", "_")
    if key == "exchange_rate":
        return "Exchange Rate"
    return text


def normalize_horizon(value) -> str:
    try:
        return str(int(float(value)))
    except (TypeError, ValueError):
        return str(value).strip()


def detail_key(model, dataset, horizon) -> tuple[str, str, str]:
    return (str(model).strip(), normalize_dataset(dataset), normalize_horizon(horizon))


def to_float(value):
    if value in ("", None):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def load_scheme6_rows(path: Path, variant: str) -> tuple[dict[tuple[str, str, str], dict], list[dict]]:
    selected: dict[tuple[str, str, str], dict] = {}
    all_rows: list[dict] = []
    if not path.exists():
        return selected, all_rows

    with path.open(newline="", encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            if row.get("family") != "COSA_SCHEME6":
                continue
            all_rows.append(row)
            if row.get("variant") != variant:
                continue
            key = detail_key(row.get("model", ""), row.get("dataset", ""), row.get("horizon", ""))
            selected[key] = row
    return selected, all_rows


def cell_style_from(source, target) -> None:
    if source.has_style:
        target._style = copy(source._style)
    if source.number_format:
        target.number_format = source.number_format
    if source.alignment:
        target.alignment = copy(source.alignment)


def header_map(ws) -> dict[str, int]:
    return {
        str(ws.cell(1, col).value).strip(): col
        for col in range(1, ws.max_column + 1)
        if ws.cell(1, col).value not in ("", None)
    }


def ensure_detail_columns(ws, method_label: str) -> dict[str, int]:
    headers = [
        ("mse", f"{method_label} MSE"),
        ("base_abs", f"{method_label} vs Baseline Abs Drop"),
        ("base_pct", f"{method_label} vs Baseline %"),
        ("f_abs", f"{method_label} vs COSA-F Abs Drop"),
        ("f_pct", f"{method_label} vs COSA-F %"),
        ("p_abs", f"{method_label} vs COSA-P Abs Drop"),
        ("p_pct", f"{method_label} vs COSA-P %"),
        ("spectral_abs", f"{method_label} vs Spectral Abs Drop"),
        ("spectral_pct", f"{method_label} vs Spectral %"),
        ("seeds", f"{method_label} Seeds"),
        ("missing", f"{method_label} Missing Seeds"),
        ("variant", f"{method_label} Variant"),
    ]
    existing = header_map(ws)
    result: dict[str, int] = {}
    source_header = ws.cell(1, existing.get("Spectral MSE", 7))
    source_body = ws.cell(2, existing.get("Spectral MSE", 7))

    for key, title in headers:
        col = existing.get(title)
        if col is None:
            col = ws.max_column + 1
            ws.cell(1, col).value = title
        result[key] = col
        cell_style_from(source_header, ws.cell(1, col))
        for row in range(2, ws.max_row + 1):
            cell_style_from(source_body, ws.cell(row, col))
        if key.endswith("_pct"):
            for row in range(2, ws.max_row + 1):
                ws.cell(row, col).number_format = "0.00%"
        elif key.endswith("_abs") or key == "mse":
            for row in range(2, ws.max_row + 1):
                ws.cell(row, col).number_format = "0.000000"
        else:
            for row in range(2, ws.max_row + 1):
                ws.cell(row, col).number_format = "@"

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    return result


def seed_label(row: dict) -> str:
    seeds = str(row.get("seeds", "")).strip()
    count = str(row.get("seed_count", "")).strip()
    if not seeds:
        return ""
    return f"{count or len(seeds.split())} seeds: {seeds}"


def fill_detail_sheet(ws, rows_by_key: dict, variant: str, method_label: str) -> tuple[int, int, dict[str, int]]:
    existing = header_map(ws)
    cols = ensure_detail_columns(ws, method_label)
    max_row = ws.max_row

    source_cols = {
        "Baseline": existing["Baseline MSE"],
        "COSA-F": existing["COSA-F MSE"],
        "COSA-P": existing["COSA-P MSE"],
        "Spectral": existing["Spectral MSE"],
    }
    abs_cols = {
        "Baseline": cols["base_abs"],
        "COSA-F": cols["f_abs"],
        "COSA-P": cols["p_abs"],
        "Spectral": cols["spectral_abs"],
    }
    pct_cols = {
        "Baseline": cols["base_pct"],
        "COSA-F": cols["f_pct"],
        "COSA-P": cols["p_pct"],
        "Spectral": cols["spectral_pct"],
    }

    filled = 0
    incomplete = 0
    for row_idx in range(2, max_row + 1):
        key = detail_key(ws.cell(row_idx, 1).value, ws.cell(row_idx, 2).value, ws.cell(row_idx, 3).value)
        scheme6 = rows_by_key.get(key)
        mse_cell = ws.cell(row_idx, cols["mse"])
        if scheme6:
            mse = to_float(scheme6.get("test_mse_mean", ""))
            mse_cell.value = mse if mse is not None else None
            ws.cell(row_idx, cols["seeds"]).value = seed_label(scheme6)
            ws.cell(row_idx, cols["missing"]).value = scheme6.get("missing_seeds", "")
            ws.cell(row_idx, cols["variant"]).value = scheme6.get("variant", variant)
            filled += 1
            if scheme6.get("missing_seeds", ""):
                incomplete += 1
        else:
            mse_cell.value = None
            ws.cell(row_idx, cols["seeds"]).value = ""
            ws.cell(row_idx, cols["missing"]).value = ""
            ws.cell(row_idx, cols["variant"]).value = variant

        mse_ref = f"{get_column_letter(cols['mse'])}{row_idx}"
        for name, source_col in source_cols.items():
            source_ref = f"{get_column_letter(source_col)}{row_idx}"
            abs_ref = f"{get_column_letter(abs_cols[name])}{row_idx}"
            ws.cell(row_idx, abs_cols[name]).value = f'=IF(OR({source_ref}="",{mse_ref}=""),"",{source_ref}-{mse_ref})'
            ws.cell(row_idx, pct_cols[name]).value = f'=IFERROR({abs_ref}/{source_ref},"")'

    return filled, incomplete, cols


def style_summary_row(ws, row_idx: int) -> None:
    for col in range(1, 11):
        cell_style_from(ws.cell(5, col), ws.cell(row_idx, col))


def write_summary(ws, detail_max_row: int, cols: dict[str, int], method_label: str) -> None:
    specs = [
        (f"{method_label} vs Baseline", "Baseline", method_label, cols["base_abs"], cols["base_pct"], f"(Baseline - {method_label}) / Baseline"),
        (f"{method_label} vs COSA-F", "COSA-F", method_label, cols["f_abs"], cols["f_pct"], f"(COSA-F - {method_label}) / COSA-F"),
        (f"{method_label} vs COSA-P", "COSA-P", method_label, cols["p_abs"], cols["p_pct"], f"(COSA-P - {method_label}) / COSA-P"),
        (f"{method_label} vs Spectral", "Spectral", method_label, cols["spectral_abs"], cols["spectral_pct"], f"(Spectral - {method_label}) / Spectral"),
    ]
    start_row = 16
    for offset, spec in enumerate(specs):
        row_idx = start_row + offset
        comparison, from_label, to_label, abs_col, pct_col, formula_text = spec
        abs_letter = get_column_letter(abs_col)
        pct_letter = get_column_letter(pct_col)
        style_summary_row(ws, row_idx)
        ws.cell(row_idx, 1).value = comparison
        ws.cell(row_idx, 2).value = from_label
        ws.cell(row_idx, 3).value = to_label
        ws.cell(row_idx, 4).value = f"=COUNT(Detail_All!{abs_letter}2:{abs_letter}{detail_max_row})"
        ws.cell(row_idx, 5).value = f"=AVERAGE(Detail_All!{abs_letter}2:{abs_letter}{detail_max_row})"
        ws.cell(row_idx, 6).value = f"=AVERAGE(Detail_All!{pct_letter}2:{pct_letter}{detail_max_row})"
        ws.cell(row_idx, 7).value = f'=COUNTIF(Detail_All!{abs_letter}2:{abs_letter}{detail_max_row},">0")'
        ws.cell(row_idx, 8).value = f'=IFERROR(G{row_idx}/D{row_idx},"")'
        ws.cell(row_idx, 9).value = f'=COUNTIF(Detail_All!{abs_letter}2:{abs_letter}{detail_max_row},"<0")'
        ws.cell(row_idx, 10).value = formula_text
        ws.cell(row_idx, 6).number_format = "0.00%"
        ws.cell(row_idx, 8).number_format = "0.00%"


def ensure_group_columns(ws, specs: list[tuple[str, str, int]], criterion_col: str, detail_max_row: int) -> None:
    existing = header_map(ws)
    start_col = ws.max_column + 1
    for idx, (title, detail_col_letter, abs_col) in enumerate(specs):
        col = existing.get(title)
        if col is None:
            col = start_col
            start_col += 1
            ws.cell(1, col).value = title
        cell_style_from(ws.cell(1, 5), ws.cell(1, col))
        ws.cell(1, col).alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        for row_idx in range(2, ws.max_row + 1):
            cell_style_from(ws.cell(row_idx, 5), ws.cell(row_idx, col))
            criterion_ref = f"A{row_idx}"
            if criterion_col == "B":
                criterion_ref = f"A{row_idx}"
            if title.endswith("Win Rate"):
                ws.cell(row_idx, col).value = (
                    f'=IFERROR(COUNTIFS(Detail_All!${criterion_col}$2:${criterion_col}${detail_max_row},'
                    f'{criterion_ref},Detail_All!${get_column_letter(abs_col)}$2:${get_column_letter(abs_col)}${detail_max_row},">0")/B{row_idx},"")'
                )
            else:
                ws.cell(row_idx, col).value = (
                    f'=AVERAGEIF(Detail_All!${criterion_col}$2:${criterion_col}${detail_max_row},'
                    f'{criterion_ref},Detail_All!${detail_col_letter}$2:${detail_col_letter}${detail_max_row})'
                )
            ws.cell(row_idx, col).number_format = "0.00%"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"


def write_group_summaries(wb, detail_max_row: int, cols: dict[str, int], method_label: str) -> None:
    pct_specs = [
        (f"{method_label} vs Baseline %", get_column_letter(cols["base_pct"]), cols["base_abs"]),
        (f"{method_label} vs COSA-F %", get_column_letter(cols["f_pct"]), cols["f_abs"]),
        (f"{method_label} vs COSA-P %", get_column_letter(cols["p_pct"]), cols["p_abs"]),
        (f"{method_label} vs Spectral %", get_column_letter(cols["spectral_pct"]), cols["spectral_abs"]),
        (f"{method_label}>Baseline Win Rate", get_column_letter(cols["base_pct"]), cols["base_abs"]),
        (f"{method_label}>F Win Rate", get_column_letter(cols["f_pct"]), cols["f_abs"]),
        (f"{method_label}>P Win Rate", get_column_letter(cols["p_pct"]), cols["p_abs"]),
        (f"{method_label}>Spectral Win Rate", get_column_letter(cols["spectral_pct"]), cols["spectral_abs"]),
    ]
    ensure_group_columns(wb["By_Model"], pct_specs, "A", detail_max_row)
    ensure_group_columns(wb["By_Dataset"], pct_specs, "B", detail_max_row)


def reset_sheet(wb, name: str):
    if name in wb.sheetnames:
        del wb[name]
    return wb.create_sheet(name)


def write_comparison_sheet(
    wb,
    name: str,
    title: str,
    from_label: str,
    to_label: str,
    from_col: int,
    to_col: int,
    detail_max_row: int,
) -> None:
    ws = reset_sheet(wb, name)
    headers = ["Model", "Dataset", "Horizon", "From", "To", "From MSE", "To MSE", "Abs Drop", "% Improvement", "Improved?"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(1, col)
        cell.value = header
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.cell(1, 12).value = title
    ws.cell(1, 12).font = Font(bold=True)

    from_letter = get_column_letter(from_col)
    to_letter = get_column_letter(to_col)
    for row_idx in range(2, detail_max_row + 1):
        ws.cell(row_idx, 1).value = f"=Detail_All!A{row_idx}"
        ws.cell(row_idx, 2).value = f"=Detail_All!B{row_idx}"
        ws.cell(row_idx, 3).value = f"=Detail_All!C{row_idx}"
        ws.cell(row_idx, 4).value = from_label
        ws.cell(row_idx, 5).value = to_label
        ws.cell(row_idx, 6).value = f"=Detail_All!{from_letter}{row_idx}"
        ws.cell(row_idx, 7).value = f"=Detail_All!{to_letter}{row_idx}"
        ws.cell(row_idx, 8).value = f'=IF(OR(F{row_idx}="",G{row_idx}=""),"",F{row_idx}-G{row_idx})'
        ws.cell(row_idx, 9).value = f'=IFERROR(H{row_idx}/F{row_idx},"")'
        ws.cell(row_idx, 10).value = f'=IF(H{row_idx}>0,"Yes",IF(H{row_idx}<0,"No","Tie"))'
        ws.cell(row_idx, 6).number_format = "0.000000"
        ws.cell(row_idx, 7).number_format = "0.000000"
        ws.cell(row_idx, 8).number_format = "0.000000"
        ws.cell(row_idx, 9).number_format = "0.00%"

    ws.cell(2, 12).value = "Metric"
    ws.cell(2, 13).value = "Value"
    ws.cell(3, 12).value = "N"
    ws.cell(3, 13).value = f"=COUNT(H2:H{detail_max_row})"
    ws.cell(4, 12).value = "Avg % Improvement"
    ws.cell(4, 13).value = f"=AVERAGE(I2:I{detail_max_row})"
    ws.cell(5, 12).value = "Win Rate"
    ws.cell(5, 13).value = f'=COUNTIF(H2:H{detail_max_row},">0")/M3'
    ws.cell(4, 13).number_format = "0.00%"
    ws.cell(5, 13).number_format = "0.00%"

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:J{detail_max_row}"
    widths = {1: 16, 2: 16, 3: 10, 4: 12, 5: 12, 6: 14, 7: 14, 8: 14, 9: 14, 10: 12, 12: 24, 13: 16}
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def write_comparison_sheets(wb, detail_max_row: int, cols: dict[str, int], method_label: str) -> None:
    existing = header_map(wb["Detail_All"])
    to_col = cols["mse"]
    write_comparison_sheet(
        wb, "G5_V5_vs_Base", f"{method_label} improvement over Baseline",
        "Baseline", method_label, existing["Baseline MSE"], to_col, detail_max_row
    )
    write_comparison_sheet(
        wb, "G6_V5_vs_F", f"{method_label} improvement over COSA-F",
        "COSA-F", method_label, existing["COSA-F MSE"], to_col, detail_max_row
    )
    write_comparison_sheet(
        wb, "G7_V5_vs_P", f"{method_label} improvement over COSA-P",
        "COSA-P", method_label, existing["COSA-P MSE"], to_col, detail_max_row
    )
    write_comparison_sheet(
        wb, "G8_V5_vs_Spec", f"{method_label} improvement over Spectral",
        "Spectral", method_label, existing["Spectral MSE"], to_col, detail_max_row
    )


def write_all_variants_sheet(wb, all_rows: list[dict]) -> None:
    ws = reset_sheet(wb, "Scheme6_All")
    headers = [
        "Variant",
        "Model",
        "Dataset",
        "Horizon",
        "Seed Count",
        "Seeds",
        "Missing Seeds",
        "Test MSE Mean",
        "Test MSE Std",
        "Test MAE Mean",
        "Test MAE Std",
        "Complete?",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def sort_key(row: dict):
        return (
            row.get("variant", ""),
            row.get("model", ""),
            normalize_dataset(row.get("dataset", "")),
            normalize_horizon(row.get("horizon", "")),
        )

    for row in sorted(all_rows, key=sort_key):
        ws.append([
            row.get("variant", ""),
            row.get("model", ""),
            normalize_dataset(row.get("dataset", "")),
            normalize_horizon(row.get("horizon", "")),
            to_float(row.get("seed_count", "")),
            row.get("seeds", ""),
            row.get("missing_seeds", ""),
            to_float(row.get("test_mse_mean", "")),
            to_float(row.get("test_mse_std", "")),
            to_float(row.get("test_mae_mean", "")),
            to_float(row.get("test_mae_std", "")),
            "Yes" if not row.get("missing_seeds", "") else "No",
        ])
    for col_idx in range(1, len(headers) + 1):
        letter = get_column_letter(col_idx)
        ws.column_dimensions[letter].width = 18
    for row_idx in range(2, ws.max_row + 1):
        for col_idx in (8, 9, 10, 11):
            ws.cell(row_idx, col_idx).number_format = "0.000000"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"


def set_recalculate_on_open(wb) -> None:
    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
    except AttributeError:
        pass


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", default=DEFAULT_TEMPLATE, help="5.26 Excel workbook to use as template")
    parser.add_argument("--scheme6-averaged", default=DEFAULT_SCHEME6_AVERAGED, help="Scheme6 averaged CSV")
    parser.add_argument("--output", default=DEFAULT_OUTPUT, help="output workbook path")
    parser.add_argument("--variant", default="adapter_core", help="Scheme6 variant to populate as the main V5 column")
    parser.add_argument("--method-label", default="V5", help="label to use in workbook columns and comparisons")
    parser.add_argument("--require-complete", action="store_true", help="fail if any template row is missing the selected variant or a seed")
    args = parser.parse_args()

    template_path = project_path(args.template)
    averaged_path = project_path(args.scheme6_averaged)
    output_path = project_path(args.output)

    rows_by_key, all_rows = load_scheme6_rows(averaged_path, args.variant)
    wb = load_workbook(template_path)
    detail_ws = wb["Detail_All"]
    filled, incomplete, cols = fill_detail_sheet(detail_ws, rows_by_key, args.variant, args.method_label)
    missing_template_rows = (detail_ws.max_row - 1) - filled

    write_summary(wb["Summary"], detail_ws.max_row, cols, args.method_label)
    write_group_summaries(wb, detail_ws.max_row, cols, args.method_label)
    write_comparison_sheets(wb, detail_ws.max_row, cols, args.method_label)
    write_all_variants_sheet(wb, all_rows)
    set_recalculate_on_open(wb)

    if args.require_complete and (missing_template_rows or incomplete):
        raise SystemExit(
            f"Selected variant is incomplete: filled={filled}, "
            f"missing_template_rows={missing_template_rows}, incomplete_seed_rows={incomplete}"
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f"wrote {output_path}")
    print(
        f"selected_variant={args.variant} filled_rows={filled} "
        f"missing_template_rows={missing_template_rows} incomplete_seed_rows={incomplete} all_variant_rows={len(all_rows)}"
    )


if __name__ == "__main__":
    main()
